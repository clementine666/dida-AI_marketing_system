# -*- coding: utf-8 -*-
"""生成活动档案字段沟通 Excel（Jim 对齐用 · 四行格式）"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parents[1] / "docs" / "活动档案表字段沟通-Jim.xlsx"

ROW_LABELS = [
    "依赖流程模块",
    "对应的必须字段名称",
    "字段对系统的影响",
    "字段的定义解释",
]

# 每个 block: module_desc + list of {name, impact, definition}
BLOCKS = [
    {
        "title": "环节1 · 活动生成（上）— 评审必备：是什么、何时、从哪来、做哪里",
        "module": (
            "活动生成环节 — 评审时必须看清楚：这是什么活动、什么类型、什么时间做、活动来源、"
            "做哪些国家/城市；采纳后进全年计划池，是后续自动化与复盘的「母本」。"
        ),
        "fields": [
            ("活动来源", "缺：无法区分飞书日历/中途录入/创意，评审与溯源断档", "例：飞书营销日历、中途录入、AI创意"),
            ("活动类型", "缺：标准/创意/其他无法分流，必填规则对不上", "仅三种：人工标准 / AI创意 / 其他"),
            ("活动推广月", "缺：全年计划池、月历无法归月", "规划归属月，例：2027-09"),
            ("活动上线时间（周二）", "缺：所有任务截止日（T-7、T-14）无法计算", "对外上线日，例：2027-09-14"),
            ("月度主活动主题", "缺：同月多场无法归组、主题筛选失效", "当月统一大主题"),
            ("副活动主题（按目的地/区域）", "缺：日历无活动名、同主题多目的地无法区分", "具体场次名/活动名"),
            ("覆盖目的地/城市", "缺：资源分析、圈客选品、目的地筛选无输入", "例：新加坡；可含多城"),
            ("对应出游/离店窗口", "缺：资源准备窗口、GP测算与真实需求错位", "客户离店高峰月，例：2027-10"),
        ],
    },
    {
        "title": "环节1 · 活动生成（下）— 为什么做、怎么做、对谁做、预期价值",
        "module": "活动生成环节（续）— 支撑 Jim 复盘 ①方案解读、②原来想的是什么；没有这些，活动总结只能报数、无法对比。",
        "fields": [
            ("选目的地数据依据·节庆/旺季", "缺：评审无法论证选题；复盘无法对照立项假设", "Top/峰值/淡旺季文字依据"),
            ("立项背景摘要", "缺：没有活动总结的故事线", "3–5句：为什么现在做"),
            ("核心目标", "缺：无法判断活动成不成功", "一句话总目标"),
            ("TTV 目标", "缺：达成率、T+1推送、复盘对比无基准", "例：100–200万"),
            ("订单目标", "缺：转化达成率无法计算", "例：500+"),
            ("GP 目标", "缺：GP复盘与风控无预期", "GP目标值或区间"),
            ("活动玩法/机制", "缺：方案任务、运营brief无法生成", "券/满减/专题等可执行玩法"),
            ("素材需求清单", "缺：素材线任务无法自动拆分", "Banner,海报,文案（逗号分隔，一项一条任务）"),
            ("目标客户描述", "缺：资源线圈客不知道圈谁", "圈选逻辑、行为特征"),
            ("酒店选品策略", "缺：资源线选品不知道选什么", "例：核心商圈高星优先"),
            ("上线位置", "缺：Jim复盘①无法写清资源位", "例：官网Banner+活动专区"),
            ("交付形式", "缺：审核不知验收Banner还是券", "Banner/Coupon/专题页/组合活动"),
        ],
    },
    {
        "title": "环节2 · 活动任务管理 — 自动创建、发送、催办、交付",
        "module": (
            "活动任务管理环节 — 以活动上线日为锚点，自动拆资源/方案/素材三线任务；"
            "支持飞书发送、到期催办、交付验收。原活动表无具体任务信息则无法自动化。"
        ),
        "fields": [
            ("活动上线日（系统锚点）", "缺：整板任务不生成，催办SLA全部失效", "由「活动上线时间」自动同步"),
            ("活动推广起止", "缺：监控拉数窗口、自动下线时间不明", "例：2027-09-14 ~ 2027-10-28"),
            ("资源对接时间", "缺：资源线启动时点不明", "Rachel确认节点"),
            ("资源交付时间（T-7）", "缺：萧蓉线deadline无法对齐、无法催办", "上线前7天交资源包"),
            ("物料完成时间（T-14）", "缺：素材deadline无法对齐", "上线前14天物料齐套"),
            ("资源线任务包", "缺：QBI分析/圈客/选品/GP等无人认领", "系统自动生成A–F任务"),
            ("方案线任务包", "缺：玩法/位置/券/监控确认无任务", "系统自动生成P1–P5任务"),
            ("素材线任务包", "缺：Banner/海报/文案无人制作", "由素材清单一项一条自动生成"),
            ("任务负责人", "缺：通知无法@到人", "例：萧蓉/JIM/梓淮"),
            ("任务交付说明", "缺：监督无法验收「做了什么」", "结论摘要、表格说明"),
            ("任务交付链接/附件", "缺：交付物无法回看、审核无法点开", "飞书文档、预览图链接"),
        ],
    },
    {
        "title": "环节3 · 上线审核 — T-1联合验收、过闸上线",
        "module": "上线审核环节 — 资源包+素材+配置齐套后过闸；把原来口头/群里的「定版信息」结构化，系统才能做Checklist。",
        "fields": [
            ("展示策略", "缺：运营/前端规则不一致，审核无法目视对照", "轮播位次、专区展示规则"),
            ("优惠券策略", "缺：券类活动规则未确认，存在配置风险", "满减规则；纯Banner可空"),
            ("锁定客户 ID", "缺：不能推送触达，审核不过闸", "萧蓉T-7前确认的客户清单"),
            ("锁定酒店 ID", "缺：GP风控无清单，资源包无法定版", "确认版hotel_id列表"),
            ("落地页 URL", "缺：Checklist红灯，监控无法关联页面", "活动专题完整URL"),
            ("埋点/监控事件", "缺：上线后看板全空，漏斗拉不到", "例：primary_banner_click"),
            ("活动推广下线日", "缺：无法按时自动下线", "推广结束日，由推广起止同步"),
        ],
    },
    {
        "title": "环节4 · 上线监控 — 执行期看板、T+1、异常（支撑Jim复盘③④）",
        "module": "上线监控环节 — 执行期拉数、看板、T+1推送；需要事先定「监控什么」，事中回写「实际发生什么」。",
        "fields": [
            ("过程监控指标", "缺：看板不知展示CTR/CVR等维度", "例：曝光/点击/CTR/CVR/TTV/GP"),
            ("成功标准", "缺：无法自动判断Pass/Fail", "例：TTV达标且CTR≥基线1.1x"),
            ("Banner曝光/点击（实绩）", "缺：漏斗上层断裂", "执行期数仓回写，通常不需手填"),
            ("订单/TTV/GP（实绩）", "缺：T+1总结无数可报、无法对比目标", "执行期累计值"),
            ("监控快照时间", "缺：不知数据是否过期", "最近一次拉数时间"),
            ("异常/诊断说明", "缺：异常只能人工发现", "监控Agent一句话结论"),
        ],
    },
    {
        "title": "环节5 · 复盘归档 — Jim复盘⑤⑥⑦（支撑活动总结与反哺）",
        "module": (
            "复盘归档环节 — 活动结束后总结、沉淀经验、产出有产酒店/机构清单；"
            "⑤结论⑥经验⑦清单 依赖环节1–4事先填好的方案与目标。"
        ),
        "fields": [
            ("活动复盘时间", "缺：复盘任务无法自动提醒", "建议结束后7天"),
            ("实际 TTV/GP/订单", "缺：无法与立项目标对比", "复盘定稿实绩"),
            ("目标达成率", "缺：档案库无法按效果检索", "例：TTV达成112%"),
            ("效果定性（成功/失败/部分）", "缺：Jim复盘⑤无法输出结论", "人工判定"),
            ("根因分析", "缺：只能报数，无法归因", "供给/客群/素材/节奏等"),
            ("人工复盘结论", "缺：归档缺「人话结论」", "运营师最终结论"),
            ("可复用经验", "缺：下一场活动检索不到经验", "正向：玩法/资源/时段等成功动作"),
            ("规避点", "缺：重复踩坑", "负向：失误/未达预期原因"),
            ("有产酒店清单（产出物）", "缺：Jim复盘⑦无法导出归因清单", "按订单/TTV排序导出"),
            ("有产机构ID清单（产出物）", "缺：无法做定向运营", "按订单/TTV排序导出"),
        ],
    },
]

JIM_MAP = [
    ("Jim复盘① 活动方案解读", "活动类型、主/副主题、玩法、上线位置、展示策略、优惠券策略、触达渠道、目标客户描述、落地页"),
    ("Jim复盘② 原来想的是什么", "核心目标、TTV/GP/订单目标、过程监控指标（含曝光/点击/转化预期）"),
    ("Jim复盘③ 结果发生了什么", "上述目标 + 埋点 + 推广期 → 对比实绩、漏斗转化、异常说明"),
    ("Jim复盘④ 活动最终效果", "TTV/GP/订单实绩；券活动另需：用券占比、用券TTV（报表导出）"),
    ("Jim复盘⑤ 输出结论", "成功标准 + 目标与实绩 → 定性、根因、下次调整方向"),
    ("Jim复盘⑥ 可复用经验", "可复用经验、规避点（复盘时填写）"),
    ("Jim复盘⑦ 产出物清单", "锁定客户/酒店ID + 订单归因 → 有产酒店清单、有产机构ID清单"),
]


def style_header_cell(cell):
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="E8EEF4")
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_blocks(ws, blocks, start_row=1):
    r = start_row
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for block in blocks:
        if block.get("title"):
            ws.cell(r, 1, block["title"]).font = Font(bold=True, size=12, color="1D4ED8")
            r += 1

        n = len(block["fields"])
        last_col = 1 + n

        for i, label in enumerate(ROW_LABELS):
            c = ws.cell(r + i, 1, label)
            style_header_cell(c)

        ws.cell(r, 2, block["module"])
        if n > 1:
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=last_col)

        for col, (name, impact, defn) in enumerate(block["fields"], start=2):
            ws.cell(r + 1, col, name)
            ws.cell(r + 2, col, impact)
            ws.cell(r + 3, col, defn)
            for ri in range(r, r + 4):
                for ci in range(1, last_col + 1):
                    cell = ws.cell(ri, ci)
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    cell.border = border

        r += 6

    return r


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "活动档案字段（五环节）"

    ws.column_dimensions["A"].width = 22
    for col in range(2, 20):
        ws.column_dimensions[get_column_letter(col)].width = 28

    ws.cell(1, 1, "道旅 · 活动档案表字段沟通稿（Jim / 营销活动负责人）").font = Font(bold=True, size=14)
    ws.merge_cells("A1:H1")
    ws.cell(2, 1, "说明：A列为行标题；每个字段占一列。请在此表上增删改字段，定稿后作为环节二「活动建档」页面设计依据。").font = Font(
        italic=True, color="666666"
    )
    ws.merge_cells("A2:H2")

    write_blocks(ws, BLOCKS, start_row=4)

    ws2 = wb.create_sheet("Jim复盘7段对照")
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 70
    ws2.cell(1, 1, "Jim复盘段落").font = Font(bold=True)
    ws2.cell(1, 2, "活动档案需事先具备的信息（字段/产出）").font = Font(bold=True)
    for i, (seg, fields) in enumerate(JIM_MAP, start=2):
        ws2.cell(i, 1, seg)
        ws2.cell(i, 2, fields)
        ws2.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws2.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(str(OUT))


if __name__ == "__main__":
    main()
